from __future__ import annotations

import asyncio
import json
import re
from io import BytesIO
from datetime import datetime

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from openpyxl import load_workbook
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.session import SessionLocal, get_db
from app.models import Asset, Job, Project, Sequence
from app.schemas.api import AudioBatchProcessCreate, AudioProcessJobCreate, JobOut, RenderItemCreate, RenderJobCreate, TTSJobCreate
from app.services.job_service import job_manager
from app.services.storage_service import safe_filename


router = APIRouter(tags=["jobs"])


def _next_tts_name(db: Session, project_id: str, requested_name: str) -> str:
    """Append a persistent, project-wide sequence number to a user supplied prefix."""
    prefix = safe_filename(requested_name, "voice")
    pattern = re.compile(rf"^{re.escape(prefix)}_(\d+)$")
    job_inputs = db.scalars(select(Job.input_json).where(Job.project_id == project_id, Job.kind == "tts")).all()
    used = []
    for input_json in job_inputs:
        match = pattern.match(str(input_json.get("name", "")))
        if match:
            used.append(int(match.group(1)))
    return f"{prefix}_{max(used, default=0) + 1}"


def _queue_tts_jobs(db: Session, project_id: str, payloads: list[dict]) -> list[Job]:
    jobs = [Job(project_id=project_id, kind="tts", status="queued", input_json=item) for item in payloads]
    db.add_all(jobs)
    db.commit()
    for job in jobs:
        db.refresh(job)
        job_manager.submit(job.id)
    return jobs


@router.post("/projects/{project_id}/tts/jobs", response_model=JobOut, status_code=202)
def create_tts_job(project_id: str, payload: TTSJobCreate, db: Session = Depends(get_db)):
    if not db.get(Project, project_id):
        raise HTTPException(404, "Project was not found")
    values = payload.model_dump()
    values["name"] = _next_tts_name(db, project_id, payload.name)
    return _queue_tts_jobs(db, project_id, [values])[0]


@router.post("/projects/{project_id}/tts/excel-jobs", response_model=list[JobOut], status_code=202)
async def create_tts_excel_jobs(
    project_id: str,
    settings: str = Form(...),
    upload: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    if not db.get(Project, project_id):
        raise HTTPException(404, "Project was not found")
    if not (upload.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(415, "Chỉ hỗ trợ file Excel .xlsx")
    try:
        base_payload = TTSJobCreate.model_validate_json(settings)
    except Exception as exc:
        raise HTTPException(422, f"Tham số tạo giọng không hợp lệ: {exc}") from exc
    group_name = (base_payload.group_name or "").strip()
    if not group_name:
        raise HTTPException(422, "Hãy đặt tên nhóm audio trước khi nhập Excel")
    existing_groups = [item.get("group_name") for item in db.scalars(select(Job.input_json).where(Job.project_id == project_id, Job.kind == "tts")).all()]
    existing_groups += [item.get("group_name") for item in db.scalars(select(Asset.metadata_json).where(Asset.project_id == project_id)).all()]
    if group_name.casefold() in {str(name).strip().casefold() for name in existing_groups if name}:
        raise HTTPException(409, "Tên nhóm đã được sử dụng. Hãy chọn tên nhóm mới.")
    base_payload.group_name = group_name
    content = await upload.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(413, "File Excel không được vượt quá 10 MB")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
    except Exception as exc:
        raise HTTPException(400, "Không thể đọc file Excel") from exc
    if not rows or len(rows[0]) < 2:
        raise HTTPException(400, "Excel phải có 2 cột STT và text")
    header = [str(value or "").strip().lower() for value in rows[0][:2]]
    if header != ["stt", "text"]:
        raise HTTPException(400, "Dòng đầu tiên phải có đúng 2 tiêu đề: STT và text")
    parsed: list[tuple[int, str]] = []
    for row_number, row in enumerate(rows[1:], start=2):
        stt, text = (list(row) + [None, None])[:2]
        if stt is None and (text is None or not str(text).strip()):
            continue
        try:
            numeric_stt = float(stt)
            number = int(numeric_stt)
            if numeric_stt != number:
                raise ValueError
        except (TypeError, ValueError):
            raise HTTPException(400, f"STT không hợp lệ tại dòng {row_number}")
        value = str(text or "").strip()
        if not value:
            raise HTTPException(400, f"Thiếu text tại dòng {row_number}")
        parsed.append((number, value))
    if not parsed:
        raise HTTPException(400, "File Excel không có nội dung để tạo giọng")
    if len(parsed) > 1000:
        raise HTTPException(400, "Mỗi lần chỉ được nhập tối đa 1000 đoạn")
    expected = list(range(1, len(parsed) + 1))
    if [number for number, _ in parsed] != expected:
        raise HTTPException(400, "Cột STT phải liên tục từ 1 đến n, không trùng hoặc bỏ số")
    prefix = safe_filename(base_payload.name, "voice")
    first_name = _next_tts_name(db, project_id, prefix)
    first_index = int(first_name.rsplit("_", 1)[1])
    payloads = []
    for offset, (number, value) in enumerate(parsed):
        item = base_payload.model_dump()
        item.update(text=value, name=f"{prefix}_{first_index + offset}", source_stt=number)
        payloads.append(item)
    return _queue_tts_jobs(db, project_id, payloads)


@router.post("/assets/{asset_id}/audio/jobs", response_model=JobOut, status_code=202)
def create_audio_process_job(
    asset_id: str, payload: AudioProcessJobCreate, db: Session = Depends(get_db)
):
    asset = db.get(Asset, asset_id)
    if not asset or asset.type not in {"tts_audio", "external_audio"}:
        raise HTTPException(404, "Audio asset was not found")
    if payload.trim_end is not None and payload.trim_end <= payload.trim_start:
        raise HTTPException(400, "Trim end must be greater than trim start")
    if asset.duration_ms and payload.trim_start >= asset.duration_ms / 1000:
        raise HTTPException(400, "Trim start is outside the audio duration")
    job = Job(
        project_id=asset.project_id,
        kind="audio_process",
        status="queued",
        input_json={"asset_id": asset.id, **payload.model_dump()},
    )
    db.add(job)
    db.commit()
    db.refresh(job)
    job_manager.submit(job.id)
    return job


@router.post("/projects/{project_id}/audio/batch-jobs", response_model=list[JobOut], status_code=202)
def create_audio_batch_jobs(
    project_id: str, payload: AudioBatchProcessCreate, db: Session = Depends(get_db)
):
    assets = db.scalars(
        select(Asset).where(Asset.id.in_(payload.asset_ids), Asset.project_id == project_id)
    ).all()
    if len(assets) != len(set(payload.asset_ids)) or any(
        asset.type not in {"tts_audio", "external_audio"} for asset in assets
    ):
        raise HTTPException(400, "Một hoặc nhiều audio không hợp lệ")
    asset_map = {asset.id: asset for asset in assets}
    jobs: list[Job] = []
    speed_label = str(payload.speed).replace(".", "-")
    for asset_id in payload.asset_ids:
        asset = asset_map[asset_id]
        job = Job(
            project_id=project_id,
            kind="audio_process",
            status="queued",
            input_json={
                "asset_id": asset.id,
                "name": f"{asset.display_name}-{speed_label}x",
                "speed": payload.speed,
                "trim_start": 0,
                "trim_end": None,
                "output_format": payload.output_format,
                "group_name": asset.metadata_json.get("group_name"),
            },
        )
        db.add(job)
        jobs.append(job)
    db.commit()
    for job in jobs:
        db.refresh(job)
        job_manager.submit(job.id)
    return jobs


@router.post("/projects/{project_id}/render/jobs", response_model=list[JobOut], status_code=202)
def create_render_jobs(project_id: str, payload: RenderJobCreate, db: Session = Depends(get_db)):
    project = db.get(Project, project_id)
    if not project:
        raise HTTPException(404, "Project was not found")
    render_items = payload.items or [
        RenderItemCreate(
            sequence_id=sequence_id,
            audio_asset_id=payload.audio_asset_id,
            source_audio_policy=payload.source_audio_policy,
        )
        for sequence_id in payload.sequence_ids
    ]
    if not render_items:
        raise HTTPException(400, "At least one render item is required")
    if not payload.allow_missing_audio and any(not item.audio_asset_id for item in render_items):
        raise HTTPException(400, "Cần gán đủ audio hoặc bật cho phép xuất video thiếu audio")
    voice_ids = [item.audio_asset_id for item in render_items if item.audio_asset_id]
    if not payload.allow_duplicate_audio and len(voice_ids) != len(set(voice_ids)):
        raise HTTPException(400, "Audio bị trùng. Cần cho phép dùng trùng trước khi render.")
    sequence_ids = [item.sequence_id for item in render_items]
    sequences = db.scalars(
        select(Sequence).where(Sequence.id.in_(sequence_ids), Sequence.project_id == project_id)
    ).all()
    if len(sequences) != len(set(sequence_ids)):
        raise HTTPException(400, "One or more sequences are invalid")
    all_audio_ids = {item.audio_asset_id for item in render_items if item.audio_asset_id} | {
        item.bgm_asset_id for item in render_items if item.bgm_asset_id
    }
    if all_audio_ids:
        audios = db.scalars(
            select(Asset).where(Asset.id.in_(all_audio_ids), Asset.project_id == project_id)
        ).all()
        if len(audios) != len(all_audio_ids) or any(
            audio.type not in {"tts_audio", "external_audio"} for audio in audios
        ):
            raise HTTPException(400, "One or more audio assets are invalid")

    values = payload.model_dump()
    option_keys = {
        "duration_policy",
        "source_audio_policy",
        "resolution",
        "fps",
        "video_encoder",
        "fit_mode",
        "voice_volume",
        "source_volume",
        "bgm_volume",
    }
    base_options = {key: values[key] for key in option_keys}
    now = datetime.now()
    # One stable timestamp per click, shared by every output in this run.
    # Keep the compact format requested by the UI, e.g. 26-9-26-10-1-1.
    timestamp = f"{now.day}-{now.month}-{now.strftime('%y')}-{now.hour}-{now.minute}-{now.second}"
    group_name = payload.group_name.strip()
    if not group_name:
        raise HTTPException(400, "Tên nhóm render không được để trống")
    output_root = safe_filename(payload.output_folder, "renders")
    # One click creates one named output group, even when source batches differ.
    run_folder = f"{output_root}/{timestamp}-{safe_filename(group_name, 'nhom-render')}"
    sequence_map = {sequence.id: sequence for sequence in sequences}
    batch_indexes: dict[str, int] = {}
    jobs: list[Job] = []
    for item in render_items:
        sequence = sequence_map[item.sequence_id]
        batch_name = sequence.batch_name or "Batch"
        batch_indexes[batch_name] = batch_indexes.get(batch_name, 0) + 1
        output_index = batch_indexes[batch_name]
        options = dict(base_options)
        options["source_audio_policy"] = item.source_audio_policy
        options["voice_volume"] = item.voice_volume
        options["bgm_volume"] = item.bgm_volume
        job = Job(
            project_id=project_id,
            kind="render",
            status="queued",
            input_json={
                "sequence_id": item.sequence_id,
                "audio_asset_id": item.audio_asset_id,
                "bgm_asset_id": item.bgm_asset_id,
                "custom_name": payload.run_name,
                "output_filename": (
                    f"{timestamp}-{safe_filename(payload.run_name, 'output')}-{output_index}.mp4"
                ),
                "output_index": output_index,
                "batch_name": batch_name,
                "group_name": group_name,
                "run_timestamp": timestamp,
                "run_folder": run_folder,
                "options": options,
            },
        )
        db.add(job)
        jobs.append(job)
    db.commit()
    for job in jobs:
        db.refresh(job)
        job_manager.submit(job.id)
    return jobs


@router.get("/jobs", response_model=list[JobOut])
def list_jobs(
    project_id: str | None = None,
    limit: int = Query(default=100, ge=1, le=500),
    db: Session = Depends(get_db),
):
    statement = select(Job)
    if project_id:
        statement = statement.where(Job.project_id == project_id)
    return db.scalars(statement.order_by(Job.created_at.desc()).limit(limit)).all()


@router.get("/jobs/{job_id}", response_model=JobOut)
def get_job(job_id: str, db: Session = Depends(get_db)):
    job = db.get(Job, job_id)
    if not job:
        raise HTTPException(404, "Job was not found")
    return job


@router.post("/jobs/{job_id}/cancel", response_model=JobOut)
def cancel_job(job_id: str, db: Session = Depends(get_db)):
    job = db.get(Job, job_id)
    if not job:
        raise HTTPException(404, "Job was not found")
    if job.status in {"completed", "failed", "cancelled"}:
        return job
    if job_manager.cancel(job_id):
        job.status = "cancelled"
        job.error_message = "Cancelled by the user"
        db.commit()
        db.refresh(job)
    return job


@router.post("/jobs/{job_id}/retry", response_model=JobOut, status_code=202)
def retry_job(job_id: str, db: Session = Depends(get_db)):
    original = db.get(Job, job_id)
    if not original:
        raise HTTPException(404, "Job was not found")
    if original.status not in {"failed", "cancelled"}:
        raise HTTPException(400, "Only failed or cancelled jobs can be retried")
    retry = Job(
        project_id=original.project_id,
        kind=original.kind,
        status="queued",
        input_json=original.input_json,
    )
    db.add(retry)
    db.commit()
    db.refresh(retry)
    job_manager.submit(retry.id)
    return retry


@router.get("/events")
async def events(project_id: str):
    async def stream():
        previous = ""
        while True:
            with SessionLocal() as db:
                jobs = db.scalars(
                    select(Job).where(Job.project_id == project_id).order_by(Job.created_at.desc()).limit(100)
                ).all()
                payload = json.dumps(
                    [
                        {
                            "id": job.id,
                            "status": job.status,
                            "progress": job.progress,
                            "error_message": job.error_message,
                            "output_asset_id": job.output_asset_id,
                        }
                        for job in jobs
                    ],
                    ensure_ascii=False,
                )
            if payload != previous:
                yield f"event: jobs\ndata: {payload}\n\n"
                previous = payload
            else:
                yield ": heartbeat\n\n"
            await asyncio.sleep(1.5)

    return StreamingResponse(stream(), media_type="text/event-stream")
